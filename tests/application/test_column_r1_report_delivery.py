"""Offline delivery from the public root; real report renderers, fake ETABS I/O."""
from dataclasses import replace
from hashlib import sha256
from importlib.util import module_from_spec, spec_from_file_location
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

import test_column_public_mixed_vertical as vertical
import tbdy_engine.product_reports.building_report_package as package
from tbdy_engine.product_reports.building_report_json import export_building_report_model_json


def _delivery():
    path = Path(__file__).resolve().parents[2] / "tools/run_column_r1_public_acceptance.py"
    spec = spec_from_file_location("_e0_acceptance_delivery", path)
    module = module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_unresolved_public_result_persists_real_deterministic_delivery(monkeypatch, tmp_path):
    harness = vertical._install(monkeypatch)
    result = vertical.project.execute_project(
        vertical.population._two_column_request(),
        verified_session=harness.setup.module._FakeSession(),
        column_design_basis=replace(vertical._basis(), story_translation_tolerance=None),
        expected_combo_policy=vertical._policy(harness.combos),
    )
    assert len(result.columns) == 2
    assert all(c.status == "UNRESOLVED" for c in result.columns)
    assert all(c.controlled_design_result is None for c in result.columns)
    model = result.building_report_model
    assert model is not None
    canonical = export_building_report_model_json(model).content
    before = {c.component_id: (c.status, c.blockers) for c in result.columns}
    owner_calls = dict(harness.counters)
    delivery = _delivery()
    first = tmp_path / "first"
    second = tmp_path / "second"
    first.mkdir()
    second.mkdir()
    built = []
    build_package = package.build_building_report_package

    def capture_package(model):
        artifact = build_package(model)
        built.append(artifact)
        return artifact

    monkeypatch.setattr(package, "build_building_report_package", capture_package)
    receipt = delivery._persist_report_package(model, first)
    repeated = delivery._persist_report_package(model, second)
    data = (first / package.DEFAULT_PACKAGE_FILENAME).read_bytes()
    assert receipt == repeated == {
        "filename": package.DEFAULT_PACKAGE_FILENAME,
        "sha256": sha256(data).hexdigest(), "size_bytes": len(data),
    }
    assert data == (second / package.DEFAULT_PACKAGE_FILENAME).read_bytes()
    with ZipFile(BytesIO(data)) as archive:
        assert tuple(archive.namelist()) == package.DEFAULT_PACKAGE_MEMBERS
        assert archive.read("building_report_model.json") == canonical
        manifest = json.loads(archive.read("manifest.json"))
        assert manifest
        for view in ("engineering", "audit"):
            assert "UNRESOLVED" in archive.read(f"{view}.html").decode()
            assert len(PdfReader(BytesIO(archive.read(f"{view}.pdf"))).pages) > 0
            workbook = load_workbook(BytesIO(archive.read(f"{view}.xlsx")))
            assert workbook.sheetnames
            workbook.close()
    assert export_building_report_model_json(model).content == canonical
    assert {c.component_id: (c.status, c.blockers) for c in result.columns} == before
    assert harness.counters == owner_calls
    monkeypatch.setattr(package, "build_building_report_package", lambda model: built[0])
    with pytest.raises(FileExistsError):
        delivery._persist_report_package(model, first)
    assert (first / package.DEFAULT_PACKAGE_FILENAME).read_bytes() == data


@pytest.mark.parametrize("failing_owner", ["build_building_report_package", "verify_building_report_package"])
def test_delivery_failure_propagates_without_persisting(monkeypatch, tmp_path, failing_owner):
    def fail(*args):
        raise package.ReportPackageIntegrityError("bounded delivery failure")

    monkeypatch.setattr(package, "build_building_report_package", lambda model: object())
    monkeypatch.setattr(package, failing_owner, fail)
    with pytest.raises(package.ReportPackageIntegrityError, match="bounded delivery failure"):
        _delivery()._persist_report_package(object(), tmp_path)
    assert list(tmp_path.iterdir()) == []


def test_delivery_rejects_absent_a40_model(tmp_path):
    with pytest.raises(TypeError, match="BuildingReportModel"):
        _delivery()._persist_report_package(None, tmp_path)
    assert list(tmp_path.iterdir()) == []
